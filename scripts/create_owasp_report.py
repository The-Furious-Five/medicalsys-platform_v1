"""Genera docs/seguridad/INF-06-Top-Ten-OWASP.docx (seccion 11 del informe, formato de la Actividad 2).

Requiere python-docx. Uso: python scripts/create_owasp_report.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "seguridad" / "INF-06-Top-Ten-OWASP.docx"

INTRO = (
    "Para esta sección tomamos las diez categorías del OWASP Top 10 en su edición 2025 y revisamos, una por "
    "una, si el riesgo se presenta en MedicalSys y qué estamos haciendo para prevenirlo. No nos limitamos a "
    "describir buenas prácticas: en cada punto indicamos lo que encontramos en el sistema, lo que ya quedó "
    "corregido en este sprint y lo que todavía está pendiente. En cuatro categorías (A01, A02, A07 y A08) "
    "repetimos la misma prueba antes y después de la corrección, y esos resultados se mencionan en el texto. "
    "El detalle de los comandos y sus salidas está guardado en el repositorio, en la carpeta "
    "docs/seguridad/evidencias."
)

ITEMS = [
    (
        "1. A01:2025 – Pérdida de Control de Acceso (Broken Access Control)",
        "Sí, y era el problema más serio que tenía la plataforma. El gateway valida el token y le pasa a cada "
        "microservicio la identidad del usuario en unas cabeceras (usuario, organización, roles y permisos), y "
        "los servicios confían en ellas. El diseño es razonable, pero el archivo de Docker Compose publicaba en "
        "el equipo los puertos de los trece microservicios, del 8081 al 8093. Eso significaba que cualquiera "
        "podía llamar directamente a un servicio y escribir esas cabeceras a mano. Lo comprobamos sobre la rama "
        "develop: una petición al puerto 8088 con un rol de médico inventado, sin ningún token, devolvió los "
        "encuentros clínicos de los pacientes. Además, varios servicios, como profesionales y catálogos, no "
        "exigían ningún permiso en sus endpoints.",
        "Cerramos los puertos de todos los microservicios y de las bases de datos, Redis, Kafka y MinIO, de modo "
        "que solo el gateway queda publicado y el resto se comunica por una red interna de Docker. Repetimos la "
        "misma petición después del cambio y ya no logra conectarse. El gateway descarta cualquier cabecera de "
        "identidad que envíe el cliente y las vuelve a generar a partir del token validado, por lo que la "
        "organización sale siempre del JWT y nunca de un parámetro que el usuario pueda cambiar. Como segunda "
        "barrera, cada endpoint declara el permiso que necesita y lo que no está declarado se deniega por "
        "defecto; ese trabajo corresponde a las tareas RBAC-03 y RBAC-04, que el equipo ya concluyó. Los roles y "
        "sus permisos se administran además desde la propia aplicación, sin tocar el código ni la base de datos "
        "(RBAC-02 y RBAC-05, también concluidas). Queda pendiente la prueba cruzada entre dos organizaciones "
        "distintas con datos reales de cada una.",
    ),
    (
        "2. A02:2025 – Configuración de Seguridad Incorrecta (Security Misconfiguration)",
        "Sí. Encontramos varios valores por defecto pensados para que la demostración funcionara rápido y no "
        "para ser seguros. El CORS del gateway aceptaba cualquier origen, las rutas internas del servicio de "
        "identidad estaban abiertas con permitAll, Actuator mostraba el detalle de salud con el motor de base "
        "de datos y la versión de Redis, y las cuentas de demostración, cuya contraseña está escrita en el "
        "repositorio, se creaban solas en cada arranque. A esto se sumaban los puertos expuestos que "
        "describimos en el punto anterior.",
        "Cambiamos los valores por defecto para que el comportamiento seguro sea el que se obtiene sin "
        "configurar nada. El CORS ahora solo acepta la lista exacta de orígenes del frontend, y al probar con "
        "un dominio ajeno el gateway responde 403. Las rutas internas de identidad pasaron a estar denegadas, "
        "Actuator muestra únicamente el estado general y las cuentas de demostración vienen apagadas: quien las "
        "necesite para una presentación tiene que encenderlas a propósito. También agregamos al gateway las "
        "cabeceras de seguridad habituales (HSTS, X-Content-Type-Options, Content-Security-Policy con "
        "frame-ancestors, Referrer-Policy y Cross-Origin-Resource-Policy) y dejamos de anunciar el servidor. Lo "
        "que falta es separar formalmente los perfiles de desarrollo y producción, porque hoy la diferencia "
        "depende solo de las variables de entorno.",
    ),
    (
        "3. A03:2025 – Fallas en la Cadena de Suministro de Software (Software Supply Chain Failures)",
        "Sí. Los catorce servicios dependen de Spring Boot y de decenas de librerías que llegan de forma "
        "transitiva, además de la imagen base de Ubuntu con la que se construyen los contenedores. El escaneo "
        "inicial lo confirmó: Trivy encontró 43 hallazgos críticos y 96 altos repartidos en las catorce "
        "imágenes. Casi todos venían de cinco componentes: Tomcat, Jackson, Bouncy Castle, lz4-java y la "
        "librería libssl del sistema operativo.",
        "Actualizamos cada componente a la versión corregida más cercana, fijando la versión de forma explícita "
        "en los pom.xml, y los Dockerfile aplican las actualizaciones de seguridad de Ubuntu al construir la "
        "imagen. Después repetimos el escaneo con las mismas herramientas y opciones, y el resultado fue de "
        "cero hallazgos críticos y cero altos. El proceso quedó escrito en el procedimiento PR-04 y se ejecuta "
        "con un solo script, de manera que cualquier integrante del equipo puede repetirlo. Lo que todavía no "
        "tenemos es el escaneo dentro del pipeline de integración continua: por ahora se corre a mano antes de "
        "cada entrega, y tampoco firmamos las imágenes. El reporte completo está en la sección 13.",
    ),
    (
        "4. A04:2025 – Fallas Criptográficas (Cryptographic Failures)",
        "Sí, porque el sistema guarda información clínica y contraseñas. Lo que ya está resuelto es lo "
        "siguiente: las contraseñas se almacenan con Argon2id, la narrativa del historial se cifra con "
        "AES-256-GCM usando una clave distinta por organización, las versiones de esas claves se guardan en la "
        "base de datos del servicio de criptografía y la conexión con MySQL exige TLS. La debilidad está en el "
        "origen de las claves. Todas se derivan de una clave maestra que llega por variable de entorno y que en "
        "el MVP tiene un valor de demostración, y los tokens se firman con un secreto compartido entre "
        "identidad y el gateway.",
        "Mantenemos las claves fuera del código y el gateway se niega a arrancar si el secreto de firma tiene "
        "menos de 32 bytes. Para el despliegue, el HTTPS lo termina el túnel de Cloudflare y el gateway ya "
        "envía la cabecera HSTS. Lo pendiente es importante y preferimos decirlo con claridad: reemplazar la "
        "clave maestra de demostración por una administrada en una bóveda, definir su rotación y su respaldo, y "
        "cifrar la comunicación entre los servicios dentro de la red interna, que hoy viaja por HTTP. El "
        "inventario criptográfico, la política de gestión de claves y el gestor de contraseñas del equipo "
        "(tareas CRY-01, PO-05 y PWM-01) todavía no se han completado. Mientras eso no esté hecho, no "
        "consideramos que el cifrado sea apto para datos reales de pacientes.",
    ),
    (
        "5. A05:2025 – Inyección (Injection)",
        "Aplica como riesgo que hay que controlar, pero no encontramos una inyección explotable. Hay varias "
        "búsquedas que reciben texto libre, como las de pacientes, diagnósticos CIE-10 y medicamentos. Al "
        "revisar el código vimos que los servicios consultan la base de datos con Spring Data JPA y consultas "
        "con parámetros, y que el servicio de notificaciones, que usa acceso reactivo, enlaza los valores en "
        "lugar de concatenarlos. No hallamos consultas SQL armadas uniendo cadenas con datos del usuario.",
        "La prevención descansa en dos cosas que ya están en el código: las consultas parametrizadas y la "
        "validación de los datos de entrada en los controladores, con límites de longitud y formato en los "
        "campos. La API responde en JSON y el frontend en Angular escapa el contenido al mostrarlo, lo que "
        "reduce el riesgo de XSS. Tenemos que ser honestos con el alcance de la verificación: el escaneo con "
        "ZAP fue pasivo y no envió cargas de inyección, así que esta conclusión sale de la revisión del código "
        "y no de una prueba de ataque. Un escaneo activo sobre los endpoints de búsqueda queda como trabajo "
        "pendiente.",
    ),
    (
        "6. A06:2025 – Diseño Inseguro (Insecure Design)",
        "Sí. Varios de los problemas anteriores no eran errores de programación sino decisiones de diseño. El "
        "principal es que los microservicios confían por completo en las cabeceras que reciben, de modo que "
        "toda la seguridad depende de que nadie más que el gateway pueda hablarles. Otro es que los servicios "
        "que usan MySQL comparten un mismo usuario de base de datos, así que un servicio comprometido podría "
        "leer las tablas de otro. El cierre de sesión tampoco estaba pensado para toda la plataforma, como "
        "explicamos en el punto 7.",
        "Reforzamos la frontera de confianza cerrando la red, de forma que el supuesto en el que se apoya el "
        "diseño ahora sí se cumple, y la plataforma deniega por defecto lo que no tiene un permiso declarado. "
        "El diseño ya tenía aciertos que conservamos: cada servicio tiene su propia base de datos, el servicio "
        "de marketing no tiene ruta ni credenciales hacia el historial clínico y el servicio de criptografía "
        "solo atiende a quien presenta un secreto compartido. Queda pendiente crear un usuario de base de datos "
        "por microservicio con permisos limitados a su propia base, que es la tarea HARD-05, y autenticar las "
        "llamadas entre servicios para que la red interna no sea la única defensa.",
    ),
    (
        "7. A07:2025 – Fallas de Autenticación (Authentication Failures)",
        "Sí. El fallo concreto era que cerrar sesión no servía fuera del servicio de identidad. Identidad "
        "borraba la sesión en Redis, pero el gateway solo revisaba la firma y la fecha de expiración del token, "
        "así que el mismo token seguía funcionando hasta por una hora. Lo comprobamos sobre develop: después "
        "del cierre de sesión, el token todavía devolvía datos de historial y de pacientes. Es justo el caso de "
        "un token robado o de un empleado al que se le acaba de retirar el acceso.",
        "Ahora el gateway consulta en Redis, en cada petición, si la sesión del token sigue abierta. Si no lo "
        "está responde 401, y si Redis no contesta responde 503 en lugar de dejar pasar la petición. Repetimos "
        "la prueba con el mismo token y, tras cerrar sesión, tanto historial como pacientes devuelven 401. "
        "Además verificamos el bloqueo de cuenta que ya existía: al quinto intento fallido la cuenta queda "
        "bloqueada quince minutos, incluso si después se escribe la contraseña correcta. Quedan dos pendientes. "
        "El gateway aún no valida el emisor del token, solo la firma y la expiración. Y la política de "
        "contraseñas completa, con el bloqueo al tercer intento que pide el proyecto, junto con el inicio de "
        "sesión delegado a Keycloak, sigue en curso en las tareas IAM-02 e IAM-03; cuando se integren, el "
        "bloqueo que describimos aquí será reemplazado por el de Keycloak.",
    ),
    (
        "8. A08:2025 – Fallas en la Integridad del Software o de los Datos (Software or Data Integrity Failures)",
        "Sí, y encontramos un caso real. El webhook de WhatsApp recibe de Meta los estados de entrega y los "
        "mensajes de los pacientes, y a partir de ellos el sistema actualiza notificaciones e incluso registra "
        "la baja de un paciente de las comunicaciones. La autenticidad de esos mensajes se verifica con una "
        "firma HMAC, pero el servicio se saltaba la verificación cuando no había un secreto configurado, que "
        "era la situación por defecto. En develop, un mensaje inventado y sin firma fue aceptado con un 200.",
        "El webhook ahora rechaza con 401 cualquier mensaje que no traiga una firma válida. Si no hay secreto "
        "configurado los rechaza todos, y si el proveedor es Meta y falta el secreto, el servicio directamente "
        "no arranca. Lo probamos con un mensaje sin firma, con una firma incorrecta y con una correcta, y solo "
        "el último fue aceptado. Para los documentos clínicos ya existía un control que mantuvimos: se calcula "
        "el hash SHA-256 de cada archivo al subirlo y se vuelve a calcular en cada descarga, y si no coincide "
        "la descarga falla. Del lado del software, los cambios entran por pull request y pasan por el pipeline "
        "de integración continua. Falta verificar la integridad de las imágenes de contenedor antes de "
        "desplegarlas, y la demostración formal del hash de documentos y los checksums de los artefactos "
        "(tareas CRY-02 y CRY-03) aún no se ha realizado.",
    ),
    (
        "9. A09:2025 – Fallas en el Registro, Alerta y Monitoreo de Seguridad (Security Logging & Alerting Failures)",
        "Sí, y es la categoría donde el sistema está más atrasado. Existe un servicio de auditoría con una "
        "bitácora que distingue entre operaciones exitosas, denegadas y fallidas, y todas las peticiones llevan "
        "un identificador de correlación que permite seguirlas de un servicio a otro. El problema es la "
        "cobertura: hoy solo el servicio de pacientes envía eventos a esa bitácora. Las consultas al historial "
        "clínico, que son las más sensibles, no dejan registro de quién las hizo, y no hay ningún mecanismo de "
        "alerta.",
        "Lo que ya funciona es la base: la bitácora, su consulta por API con permisos y el identificador de "
        "correlación. La cobertura todavía no se amplió: el registro de los eventos de autenticación, de altas "
        "y bajas de usuarios, de roles y de contraseñas, los logs de aplicación seguros y la pantalla de "
        "consulta corresponden a las tareas LOG-01 a LOG-04, que a la fecha de este informe no se han "
        "iniciado. Para considerar cubierta esta categoría hace falta que el "
        "historial y los documentos registren cada lectura y cada modificación con usuario, organización, "
        "recurso, fecha y resultado, y que exista al menos una alerta ante intentos repetidos de acceso "
        "denegado. También cuidamos que en los registros no aparezcan tokens ni contraseñas; por eso la "
        "evidencia de nuestras pruebas se guardó sin los tokens.",
    ),
    (
        "10. A10:2025 – Manejo Inadecuado de Condiciones Excepcionales (Mishandling of Exceptional Conditions)",
        "Sí, como riesgo propio de una arquitectura con muchos servicios, donde un error mal manejado puede "
        "terminar revelando información interna o dejando pasar algo que debía rechazarse. Encontramos un "
        "ejemplo claro en el webhook: cuando faltaba el secreto, el servicio trataba esa situación anómala como "
        "si todo estuviera bien y aceptaba el mensaje. El detalle de salud de Actuator, que mostraba datos de "
        "la infraestructura, también entra en esta categoría.",
        "El criterio que aplicamos es que, ante la duda, el sistema rechaza. Si el gateway no puede comprobar "
        "la sesión porque Redis no responde, devuelve 503 y no deja pasar el token; si al webhook le falta el "
        "secreto, rechaza el mensaje; y si a un servicio le falta una configuración de la que depende un "
        "control de seguridad, no arranca. Los trece microservicios tienen un manejador central de excepciones "
        "que devuelve las respuestas de error con un formato uniforme. Durante las pruebas notamos un detalle "
        "que queda por corregir: el servicio de catálogos respondió 500 ante una ruta que no existe, cuando lo "
        "correcto es un 404.",
    ),
]

CONCLUSION = (
    "Las diez categorías aplican a MedicalSys en mayor o menor medida, lo cual es esperable en una plataforma "
    "que maneja información médica y está repartida en tantos servicios. En este sprint corregimos y "
    "comprobamos con pruebas los problemas de control de acceso, configuración, autenticación e integridad, y "
    "dejamos en cero los hallazgos críticos y altos de las dependencias. Los puntos más débiles que quedan son "
    "el registro de auditoría sobre el historial clínico y la gestión de la clave maestra de cifrado, y ambos "
    "deberían resolverse antes de trabajar con datos reales de pacientes."
)


def build():
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name, size in (("Title", 20), ("Heading 1", 13)):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

    def paragraph(text, label=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if label:
            p.add_run(label + " ").bold = True
        p.add_run(text)

    doc.add_paragraph("Sección 11. Top Ten OWASP", style="Title")
    paragraph(INTRO)
    for title, applies, prevention in ITEMS:
        doc.add_heading(title, level=1)
        paragraph(applies, "¿Aplica en el sistema?:")
        paragraph(prevention, "¿Cómo lo estamos previniendo?:")
    doc.add_heading("Conclusión", level=1)
    paragraph(CONCLUSION)
    doc.save(OUT)
    print(f"{OUT.name} generado ({len(ITEMS)} categorias)")


if __name__ == "__main__":
    build()
