"""Genera PO-07.docx y PR-04.docx a partir de sus fuentes Markdown en docs/seguridad/politicas/.

El Markdown es la fuente unica: se edita el .md y se vuelve a correr este script.
Requiere python-docx (pip install python-docx).

Uso: python scripts/create_security_docs.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

DOCS = Path(__file__).resolve().parents[1] / "docs" / "seguridad" / "politicas"
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.9)
    section.left_margin = section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 12)):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(20, 63, 88)
    return doc


def add_inline(paragraph, text):
    """Negrita con **...** y codigo con `...`; el resto pasa tal cual."""
    for part in INLINE.split(text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part:
            paragraph.add_run(part)


def add_code(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)


def add_table(doc, rows):
    cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(set(c) <= set("-: ") for c in row)]
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Light Grid Accent 1"
    header = any(cells[0])
    for r, row in enumerate(cells):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            add_inline(cell.paragraphs[0], text)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9.5)
                if r == 0 and header:
                    run.bold = True
    doc.add_paragraph()


def convert(source, target):
    doc = new_document()
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"\d+\. ", line):
            add_inline(doc.add_paragraph(style="List Number"), line.split(". ", 1)[1])
        elif line.strip():
            add_inline(doc.add_paragraph(), line)
        i += 1
    doc.save(target)
    print(f"{target.relative_to(DOCS.parents[2])} generado desde {source.name}")


if __name__ == "__main__":
    for code in ("PO-07", "PR-04"):
        convert(DOCS / f"{code}.md", DOCS / f"{code}.docx")
